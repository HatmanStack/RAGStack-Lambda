"""Integration tests for text extractor registry and public API."""

import io

import pytest

from ragstack_common.text_extractors import (
    ContentSniffer,
    ExtractionResult,
    extract_text,
)
from tests.fixtures.text_extractor_samples import (
    CSV_STANDARD,
    EMAIL_SIMPLE,
    FULL_HTML_PAGE,
    JSON_SIMPLE_OBJECT,
    SIMPLE_TEXT,
    XML_SIMPLE,
)


class TestExtractTextFunction:
    """Tests for the main extract_text() function."""

    def test_extracts_txt(self):
        """Test extraction routes to TextExtractor for .txt files."""
        result = extract_text(SIMPLE_TEXT.encode(), "document.txt")

        assert isinstance(result, ExtractionResult)
        assert result.file_type == "txt"
        assert "simple text file" in result.markdown.lower()

    def test_extracts_html(self):
        """Test extraction routes to HtmlExtractor for HTML content."""
        result = extract_text(FULL_HTML_PAGE.encode(), "page.html")

        assert result.file_type == "html"
        assert "Main Heading" in result.markdown

    def test_extracts_csv(self):
        """Test extraction routes to CsvExtractor for CSV content."""
        result = extract_text(CSV_STANDARD.encode(), "data.csv")

        assert result.file_type == "csv"
        assert "name" in result.markdown.lower()

    def test_extracts_json(self):
        """Test extraction routes to JsonExtractor for JSON content."""
        result = extract_text(JSON_SIMPLE_OBJECT.encode(), "config.json")

        assert result.file_type == "json"
        assert "name" in result.markdown.lower()

    def test_extracts_xml(self):
        """Test extraction routes to XmlExtractor for XML content."""
        result = extract_text(XML_SIMPLE.encode(), "data.xml")

        assert result.file_type == "xml"
        assert "root" in result.markdown.lower()

    def test_extracts_email(self):
        """Test extraction routes to EmailExtractor for email content."""
        result = extract_text(EMAIL_SIMPLE.encode(), "message.eml")

        assert result.file_type == "eml"
        assert "sender@example.com" in result.markdown

    def test_content_sniffing_overrides_extension(self):
        """Test that content sniffing can override file extension."""
        # JSON content but with .txt extension
        result = extract_text(JSON_SIMPLE_OBJECT.encode(), "data.txt")

        # Should still detect as JSON due to content
        assert result.file_type == "json"

    def test_fallback_to_text_for_unknown(self):
        """Test fallback to text extraction for unknown types."""
        result = extract_text(b"Just plain content", "file.xyz")

        # Should fall back to txt
        assert result.file_type == "txt"

    def test_handles_binary_content(self):
        """Test handling of binary content that can't be text."""
        # Random binary that doesn't match any format
        result = extract_text(bytes(range(256)), "binary.bin")

        assert isinstance(result, ExtractionResult)


class TestExtractTextBinaryFormats:
    """Tests for binary format extraction via extract_text."""

    def _create_minimal_docx(self) -> bytes:
        """Create minimal DOCX for testing."""
        from docx import Document

        doc = Document()
        doc.add_paragraph("Test content")
        buffer = io.BytesIO()
        doc.save(buffer)
        return buffer.getvalue()

    def _create_minimal_xlsx(self) -> bytes:
        """Create minimal XLSX for testing."""
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.append(["Header"])
        ws.append(["Data"])
        buffer = io.BytesIO()
        wb.save(buffer)
        return buffer.getvalue()

    def _create_minimal_epub(self) -> bytes:
        """Create minimal EPUB for testing."""
        from ebooklib import epub

        book = epub.EpubBook()
        book.set_identifier("test-id")
        book.set_title("Test Book")
        book.set_language("en")
        chapter = epub.EpubHtml(title="Ch1", file_name="ch1.xhtml")
        chapter.content = "<html><body><p>Content</p></body></html>"
        book.add_item(chapter)
        book.add_item(epub.EpubNcx())
        book.add_item(epub.EpubNav())
        book.spine = ["nav", chapter]
        buffer = io.BytesIO()
        epub.write_epub(buffer, book)
        return buffer.getvalue()

    def test_extracts_docx(self):
        """Test extraction of DOCX via extract_text."""
        docx_bytes = self._create_minimal_docx()
        result = extract_text(docx_bytes, "document.docx")

        assert result.file_type == "docx"

    def test_extracts_xlsx(self):
        """Test extraction of XLSX via extract_text."""
        xlsx_bytes = self._create_minimal_xlsx()
        result = extract_text(xlsx_bytes, "workbook.xlsx")

        assert result.file_type == "xlsx"

    def test_extracts_epub(self):
        """Test extraction of EPUB via extract_text."""
        epub_bytes = self._create_minimal_epub()
        result = extract_text(epub_bytes, "book.epub")

        assert result.file_type == "epub"


class TestContentSnifferExport:
    """Tests for ContentSniffer export."""

    def test_sniffer_is_exported(self):
        """Test that ContentSniffer is exported from public API."""
        sniffer = ContentSniffer()
        file_type, confidence = sniffer.sniff(JSON_SIMPLE_OBJECT.encode())

        assert file_type == "json"
        assert confidence >= 0.8


class TestExtractionResultExport:
    """Tests for ExtractionResult export."""

    def test_extraction_result_is_exported(self):
        """Test that ExtractionResult is exported from public API."""
        result = ExtractionResult(
            markdown="# Test",
            file_type="txt",
            title="Test",
            word_count=1,
            structural_metadata={},
            parse_warning=None,
        )
        assert isinstance(result, ExtractionResult)


if __name__ == "__main__":
    pytest.main([__file__, "-v"])
